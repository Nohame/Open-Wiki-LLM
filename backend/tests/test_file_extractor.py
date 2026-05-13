import pytest
from io import BytesIO
from fpdf import FPDF
from docx import Document as DocxDocument
from app.services.file_extractor import extract_text


def _make_pdf(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(200, 10, text=text)
    return bytes(pdf.output())


def _make_docx(text: str) -> bytes:
    buf = BytesIO()
    doc = DocxDocument()
    doc.add_paragraph(text)
    doc.save(buf)
    return buf.getvalue()


async def test_extract_txt():
    result = await extract_text(b"Hello world", "notes.txt")
    assert result == "Hello world"


async def test_extract_md():
    result = await extract_text(b"# Titre\n\nContenu", "doc.md")
    assert result == "# Titre\n\nContenu"


async def test_extract_pdf():
    data = _make_pdf("Hello PDF")
    result = await extract_text(data, "fichier.pdf")
    assert "Hello" in result


async def test_extract_docx():
    data = _make_docx("Texte du docx")
    result = await extract_text(data, "fichier.docx")
    assert "Texte du docx" in result


async def test_extract_invalid():
    with pytest.raises(ValueError, match="non support"):
        await extract_text(b"data", "script.exe")


async def test_extract_pdf_corrupted():
    with pytest.raises(ValueError, match="PDF illisible"):
        await extract_text(b"not a pdf", "bad.pdf")


async def test_extract_docx_corrupted():
    with pytest.raises(ValueError, match="DOCX illisible"):
        await extract_text(b"not a docx", "bad.docx")
